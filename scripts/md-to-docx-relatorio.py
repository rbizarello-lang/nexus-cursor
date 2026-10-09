"""Converte RELATORIO-BETA.md e INVENTARIO-BETA.md em DOCX anotáveis."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
RULE = RGBColor(0xD1, 0xD5, 0xDB)
NOTE = RGBColor(0x6B, 0x72, 0x80)
NAVY = RGBColor(0x11, 0x18, 0x27)
HEAD3 = RGBColor(0x1E, 0x3A, 0x5F)


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_runs(paragraph, text, size=11, color=INK):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=max(9, size - 1), color=MUTED)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, italic=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def style_paragraph(p, space_after=8, space_before=0, line=1.15):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): hex_color,
        },
    )
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, val in kwargs.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val.get("val", "single"))
        el.set(qn("w:sz"), val.get("sz", "4"))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), val.get("color", "CBD5E1"))
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_page_field(paragraph):
    run1 = paragraph.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    run1._r.append(fc)
    run2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    run2._r.append(it)
    run3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "end")
    run3._r.append(fc2)
    for r in (run1, run2, run3):
        set_run_font(r, size=9, color=NOTE)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        raw = lines[i].strip()
        cells = [c.strip() for c in raw.strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", (c or "").replace(" ", "")) for c in cells if c):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i


def is_decision_table(rows):
    if not rows:
        return False
    header = " | ".join(rows[0]).lower()
    return "decisão" in header and "item" in header and "#" in header


def add_grid_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)

    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)

    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            style_paragraph(p, space_after=2, space_before=2, line=1.08)
            val = row[j] if j < len(row) else ""
            add_runs(p, val.strip(), size=9 if i else 9, color=INK if i else NAVY)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                shade_cell(cell, "E5E7EB")
    spacer = doc.add_paragraph()
    style_paragraph(spacer, space_after=6, space_before=2)


def add_note_line(container, label="Anotações"):
    p = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    if p is container:
        # already a paragraph
        pass
    style_paragraph(p, space_after=4, space_before=2, line=1.15)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10, italic=True, color=NOTE)
    r2 = p.add_run("________________________________________________")
    set_run_font(r2, size=10, color=RULE)
    return p


def add_decision_card(doc, row, headers):
    def col(*names):
        low = [h.lower() for h in headers]
        for name in names:
            for i, h in enumerate(low):
                if name in h:
                    return row[i] if i < len(row) else ""
        return ""

    num = col("#")
    item = col("item")
    tecnico = col("técnico", "tecnico")
    simples = col("linguagem")
    onde = col("onde vale")
    reversivel = col("reversível", "reversivel")
    # decisão original fica de lado: vamos redesenhar as caixas

    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)

    cell = table.cell(0, 0)
    shade_cell(cell, "F8FAFC")
    set_cell_border(
        cell,
        top={"val": "single", "sz": "8", "color": "94A3B8"},
        left={"val": "single", "sz": "8", "color": "94A3B8"},
        bottom={"val": "single", "sz": "8", "color": "94A3B8"},
        right={"val": "single", "sz": "8", "color": "94A3B8"},
    )
    cell.text = ""
    title = cell.paragraphs[0]
    title.style = doc.styles["Heading 3"]
    style_paragraph(title, space_after=4, space_before=2, line=1.1)
    title_text = f"#{num.strip()}  {item.strip()}".strip()
    add_runs(title, title_text, size=12, color=HEAD3)
    for run in title.runs:
        run.bold = True

    meta = cell.add_paragraph()
    style_paragraph(meta, space_after=6, space_before=0, line=1.1)
    bits = []
    if onde.strip() and onde.strip() != "—":
        bits.append(f"**Onde vale:** {onde.strip()}")
    if reversivel.strip() and reversivel.strip() != "—":
        bits.append(f"**Reversível:** {reversivel.strip()}")
    add_runs(meta, "   ·   ".join(bits) if bits else "", size=10, color=MUTED)

    if simples.strip():
        lab = cell.add_paragraph()
        style_paragraph(lab, space_after=2, space_before=2, line=1.1)
        r = lab.add_run("Em linguagem simples")
        set_run_font(r, size=9, bold=True, color=NOTE)
        body = cell.add_paragraph()
        style_paragraph(body, space_after=6, space_before=0, line=1.12)
        add_runs(body, simples.strip(), size=11)

    if tecnico.strip():
        lab = cell.add_paragraph()
        style_paragraph(lab, space_after=2, space_before=2, line=1.1)
        r = lab.add_run("Técnico (para conferência, se precisar)")
        set_run_font(r, size=9, bold=True, color=NOTE)
        body = cell.add_paragraph()
        style_paragraph(body, space_after=6, space_before=0, line=1.1)
        add_runs(body, tecnico.strip(), size=9, color=MUTED)

    dec = cell.add_paragraph()
    style_paragraph(dec, space_after=4, space_before=4, line=1.15)
    r = dec.add_run("Decisão:   ")
    set_run_font(r, size=11, bold=True)
    r = dec.add_run("[   ]  adotar      [   ]  não")
    set_run_font(r, size=11)

    note = cell.add_paragraph()
    style_paragraph(note, space_after=2, space_before=0, line=1.2)
    r = note.add_run("Anotações: ")
    set_run_font(r, size=10, italic=True, color=NOTE)
    r2 = note.add_run("________________________________________________")
    set_run_font(r2, size=10, color=RULE)

    spacer = doc.add_paragraph()
    style_paragraph(spacer, space_after=8, space_before=2)


def setup_doc(header_text: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = header.add_run(header_text)
    set_run_font(hr, size=9, color=NOTE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("Revisar → Novo comentário   ·   p. ")
    set_run_font(fr, size=9, color=NOTE)
    add_page_field(footer)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for level, size, before, after in (
        (1, 18, 16, 8),
        (2, 14, 14, 6),
        (3, 12, 10, 4),
    ):
        st = styles[f"Heading {level}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    return doc


def convert_markdown(src: Path, out: Path, header_text: str, notes_after_h3: bool, cards_for_decisions: bool):
    lines = src.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    doc = setup_doc(header_text)

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            style_paragraph(p, space_after=6, space_before=4)
            run = p.add_run("—" * 24)
            set_run_font(run, size=8, color=RULE)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, nxt = parse_table(lines, i)
            if cards_for_decisions and is_decision_table(rows):
                headers = rows[0]
                for row in rows[1:]:
                    add_decision_card(doc, row, headers)
            else:
                add_grid_table(doc, rows)
            i = nxt
            continue

        if stripped.startswith("# ") and not stripped.startswith("##"):
            p = doc.add_heading(stripped[2:].strip(), level=1)
            for run in p.runs:
                set_run_font(run, size=18, bold=True, color=NAVY)
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_heading(stripped[3:].strip(), level=2)
            for run in p.runs:
                set_run_font(run, size=14, bold=True, color=NAVY)
            i += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:].strip(), level=3)
            for run in p.runs:
                set_run_font(run, size=12, bold=True, color=HEAD3)
            if notes_after_h3:
                note = doc.add_paragraph()
                style_paragraph(note, space_after=8, space_before=0, line=1.2)
                r = note.add_run("Anotações: ")
                set_run_font(r, size=10, italic=True, color=NOTE)
                r2 = note.add_run("________________________________________________________________")
                set_run_font(r2, size=10, color=RULE)
            i += 1
            continue

        m_ol = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_ol:
            p = doc.add_paragraph(style="List Number")
            style_paragraph(p, space_after=4, space_before=0, line=1.15)
            add_runs(p, m_ol.group(2), size=11)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            style_paragraph(p, space_after=4, space_before=0, line=1.15)
            add_runs(p, stripped[2:], size=11)
            i += 1
            continue

        p = doc.add_paragraph()
        style_paragraph(p, space_after=8, space_before=0, line=1.15)
        add_runs(p, stripped, size=11)
        i += 1

    core = doc.core_properties
    core.author = "NEXUS"
    core.title = src.stem.replace("-", " ")
    doc.save(out)
    print(f"OK  {out}  ({out.stat().st_size:,} bytes)")


def main():
    convert_markdown(
        ROOT / "docs" / "RELATORIO-BETA.md",
        ROOT / "docs" / "RELATORIO-BETA.docx",
        "NEXUS · Relatório da Nova versão (beta) · 2.0.12",
        notes_after_h3=True,
        cards_for_decisions=False,
    )
    convert_markdown(
        ROOT / "docs" / "INVENTARIO-BETA.md",
        ROOT / "docs" / "INVENTARIO-BETA.docx",
        "NEXUS · Inventário da edição Beta · itens 1–217",
        notes_after_h3=False,
        cards_for_decisions=True,
    )


if __name__ == "__main__":
    main()
